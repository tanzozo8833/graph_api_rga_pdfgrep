"""Generate graph_grep.pptx and docs_report.docx — project documentation.

Run: venv\\Scripts\\python.exe make_docs.py
"""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt as DocxPt
from docx.shared import RGBColor as DocxRGB

# ---------- Palette ----------
BLUE = RGBColor(0x00, 0x78, 0xD4)
DARK_BLUE = RGBColor(0x00, 0x47, 0x82)
GREEN = RGBColor(0x10, 0x7C, 0x41)
PURPLE = RGBColor(0x62, 0x64, 0xA7)
TEAL = RGBColor(0x03, 0x83, 0x87)
ORANGE = RGBColor(0xD8, 0x3B, 0x01)
DARK = RGBColor(0x32, 0x31, 0x30)
LIGHT_BG = RGBColor(0xF3, 0xF2, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
YELLOW = RGBColor(0xFF, 0xB9, 0x00)
LIGHT_BLUE = RGBColor(0xDE, 0xEC, 0xF9)
GRAY = RGBColor(0x60, 0x5E, 0x5C)


# ---------- pptx helpers ----------
def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def no_line(shape):
    shape.line.fill.background()


def set_text(tf, text, *, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT, font="Segoe UI"):
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font


def add_para(tf, text, *, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT, font="Segoe UI"):
    p = tf.add_paragraph()
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def title_bar(slide, title, slide_w):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, slide_w, Inches(0.85))
    fill(bar, BLUE)
    no_line(bar)
    tf = bar.text_frame
    tf.margin_left = Inches(0.5)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(tf, title, size=24, bold=True, color=WHITE)


def textbox(slide, left, top, width, height, text, **kw):
    box = slide.shapes.add_textbox(left, top, width, height)
    set_text(box.text_frame, text, **kw)
    return box


def card(slide, left, top, width, height, header, body, color):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill(box, color)
    no_line(box)
    tf = box.text_frame
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.15)
    tf.margin_bottom = Inches(0.15)
    tf.word_wrap = True
    set_text(tf, header, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_para(tf, body, size=13, color=WHITE, align=PP_ALIGN.CENTER)
    return box


def arrow_right(slide, left, top, width, height, color=BLUE):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    fill(a, color)
    no_line(a)
    return a


def arrow_down(slide, left, top, width, height, color=BLUE):
    a = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
    fill(a, color)
    no_line(a)
    return a


# ---------- pptx generation ----------
def make_pptx(path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    SW = prs.slide_width
    SH = prs.slide_height

    # ===== 1. Title =====
    s = blank_slide(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    fill(bg, BLUE)
    no_line(bg)
    side = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.4), SH)
    fill(side, YELLOW)
    no_line(side)
    textbox(s, Inches(1), Inches(2.4), Inches(11.5), Inches(1.2),
            "Microsoft 365 Keyword-Search Agent",
            size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    textbox(s, Inches(1), Inches(3.7), Inches(11.5), Inches(0.8),
            "Question answering over OneDrive documents - no vector database required",
            size=22, color=WHITE, align=PP_ALIGN.CENTER)
    textbox(s, Inches(1), Inches(5.0), Inches(11.5), Inches(0.5),
            "Inspired by the paper 'Keyword search is all you need' (arxiv 2602.23368)",
            size=14, color=LIGHT_BG, align=PP_ALIGN.CENTER)
    textbox(s, Inches(1), Inches(6.5), Inches(11.5), Inches(0.4),
            "Microsoft Graph API  -  Azure OpenAI  -  LangChain Agent",
            size=12, color=LIGHT_BG, align=PP_ALIGN.CENTER)

    # ===== 2. Problem =====
    s = blank_slide(prs)
    title_bar(s, "Problem - traditional RAG is too complex", SW)
    items = [
        ("Chunking", "Split documents into pieces; easy to lose context between sentences/tables", BLUE),
        ("Embedding", "Costly to compute vectors for every chunk", PURPLE),
        ("Vector DB", "Must operate an additional separate service", TEAL),
        ("Exact match", "May miss product codes / rare technical terms", ORANGE),
    ]
    for i, (h, t, c) in enumerate(items):
        x = Inches(0.6 + (i % 2) * 6.3)
        y = Inches(1.4 + (i // 2) * 2.7)
        card(s, x, y, Inches(6.0), Inches(2.3), h, t, c)

    # ===== 3. Solution =====
    s = blank_slide(prs)
    title_bar(s, "Solution - Agentic Keyword Search", SW)
    textbox(s, Inches(0.7), Inches(1.2), Inches(12), Inches(0.6),
            "Let the LLM drive the search itself - no fixed pipeline",
            size=20, bold=True, color=DARK_BLUE)
    feats = [
        ("LLM generates keywords", "Synonyms, abbreviations, regex, OR-joined", BLUE),
        ("Microsoft Graph", "Full-text index already available for OneDrive", GREEN),
        ("Download file on demand", "Extract text with pdfplumber / docx / xlsx / pptx", PURPLE),
        ("Grep with context", "Multi-pattern regex, +/-N lines around each hit", TEAL),
        ("LLM retries itself", "Switches strategy if nothing is found", ORANGE),
    ]
    for i, (h, t, c) in enumerate(feats):
        x = Inches(0.6 + i * 2.5)
        y = Inches(2.5)
        card(s, x, y, Inches(2.4), Inches(2.7), h, t, c)
    textbox(s, Inches(0.7), Inches(5.7), Inches(12), Inches(0.6),
            "Microsoft already provides a keyword index - no need to rebuild it",
            size=16, color=GRAY, align=PP_ALIGN.CENTER)

    # ===== 4. Actors =====
    s = blank_slide(prs)
    title_bar(s, "Participating actors", SW)

    # User
    card(s, Inches(0.5), Inches(1.3), Inches(2.4), Inches(1.1),
         "User", "The person asking", DARK_BLUE)
    arrow_right(s, Inches(2.95), Inches(1.65), Inches(0.4), Inches(0.4))
    # Browser
    card(s, Inches(3.4), Inches(1.3), Inches(2.4), Inches(1.1),
         "Browser", "Chat panel + SSE", BLUE)
    arrow_right(s, Inches(5.85), Inches(1.65), Inches(0.4), Inches(0.4))
    # Flask
    card(s, Inches(6.3), Inches(1.3), Inches(2.4), Inches(1.1),
         "Flask", "Route /ask + MSAL", PURPLE)
    arrow_right(s, Inches(8.75), Inches(1.65), Inches(0.4), Inches(0.4))
    # Agent
    card(s, Inches(9.2), Inches(1.3), Inches(3.5), Inches(1.1),
         "LangChain Agent", "Tool-use loop", ORANGE)

    # Down arrow from agent
    arrow_down(s, Inches(10.7), Inches(2.5), Inches(0.4), Inches(0.5))

    # Three branches under agent
    card(s, Inches(2.0), Inches(3.3), Inches(3.3), Inches(1.1),
         "Azure OpenAI", "gpt-4.1-mini decides", BLUE)
    card(s, Inches(5.5), Inches(3.3), Inches(3.3), Inches(1.1),
         "Microsoft Graph", "Search & Download", GREEN)
    card(s, Inches(9.0), Inches(3.3), Inches(3.3), Inches(1.1),
         "Local Modules", "Extractors + Cache", TEAL)

    # Down arrow
    arrow_down(s, Inches(6.5), Inches(4.5), Inches(0.4), Inches(0.5))

    # Tools row
    card(s, Inches(1.5), Inches(5.2), Inches(3.0), Inches(1.4),
         "graph_search", "Find files by keyword", BLUE)
    card(s, Inches(5.0), Inches(5.2), Inches(3.0), Inches(1.4),
         "fetch_file_text", "Download + extract", GREEN)
    card(s, Inches(8.5), Inches(5.2), Inches(3.0), Inches(1.4),
         "grep_context", "Regex + context window", PURPLE)

    # ===== 5. 3 Tools =====
    s = blank_slide(prs)
    title_bar(s, "The Agent's three tools", SW)

    cols = [
        ("graph_search", BLUE,
         ["Parameter: query (string)",
          "Calls POST /search/query",
          "Returns list of files + snippets",
          "OR-joined: 'k1 OR k2 OR k3'"]),
        ("fetch_file_text", GREEN,
         ["Parameter: item_id",
          "Downloads bytes from Graph CDN",
          "Extract: PDF / DOCX / XLSX / PPTX",
          "Caches text in RAM"]),
        ("grep_context", PURPLE,
         ["Parameters: item_id, patterns, +/-N",
          "Regex IGNORECASE over the cache",
          "Merges overlapping windows",
          "Returns +/-5 lines around each hit"]),
    ]
    for i, (h, c, bullets) in enumerate(cols):
        x = Inches(0.5 + i * 4.3)
        # header
        hdr = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  x, Inches(1.2), Inches(4.0), Inches(0.9))
        fill(hdr, c)
        no_line(hdr)
        set_text(hdr.text_frame, h, size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        # body
        body = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   x, Inches(2.2), Inches(4.0), Inches(4.5))
        fill(body, LIGHT_BG)
        no_line(body)
        btf = body.text_frame
        btf.margin_left = Inches(0.2)
        btf.margin_right = Inches(0.2)
        btf.margin_top = Inches(0.2)
        btf.word_wrap = True
        set_text(btf, "- " + bullets[0], size=14, color=DARK)
        for b in bullets[1:]:
            add_para(btf, "- " + b, size=14, color=DARK)

    # ===== 6. Workflow =====
    s = blank_slide(prs)
    title_bar(s, "End-to-end workflow", SW)

    steps = [
        ("1", "User asks", "Types a question into the chat panel", BLUE),
        ("2", "Flask /ask", "Initializes the agent with the token", PURPLE),
        ("3", "LLM produces keywords", "OR-joined synonyms", ORANGE),
        ("4", "graph_search", "Microsoft Graph returns files + snippets", GREEN),
        ("5", "Enough?", "Are the snippets sufficient to answer?", TEAL),
        ("6", "fetch + grep", "Read deeper if needed", DARK_BLUE),
        ("7", "Final answer", "Synthesis + citations", BLUE),
    ]
    for i, (n, h, t, c) in enumerate(steps):
        x = Inches(0.3 + i * 1.85)
        y = Inches(1.7)
        # Number circle
        num = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.55), y, Inches(0.6), Inches(0.6))
        fill(num, c)
        no_line(num)
        set_text(num.text_frame, n, size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        # Card
        bx = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 x, Inches(2.5), Inches(1.7), Inches(2.0))
        fill(bx, LIGHT_BG)
        no_line(bx)
        btf = bx.text_frame
        btf.margin_left = Inches(0.08)
        btf.margin_right = Inches(0.08)
        btf.margin_top = Inches(0.1)
        btf.word_wrap = True
        set_text(btf, h, size=13, bold=True, color=c, align=PP_ALIGN.CENTER)
        add_para(btf, t, size=10, color=DARK, align=PP_ALIGN.CENTER)
        # Arrow except last
        if i < len(steps) - 1:
            arrow_right(s, x + Inches(1.72), y + Inches(0.18), Inches(0.13), Inches(0.25), color=GRAY)

    textbox(s, Inches(0.5), Inches(5.0), Inches(12.5), Inches(0.6),
            "The LLM decides the next step itself - this is not a fixed pipeline",
            size=18, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)
    textbox(s, Inches(0.5), Inches(5.7), Inches(12.5), Inches(0.6),
            "Every step is streamed to the browser via Server-Sent Events (SSE)",
            size=14, color=GRAY, align=PP_ALIGN.CENTER)

    # ===== 7. Search step =====
    s = blank_slide(prs)
    title_bar(s, "Search step - find files via keyword", SW)

    card(s, Inches(0.7), Inches(2.0), Inches(3.5), Inches(2.0),
         "User Query", '"What is the metrics\nfor RAG?"', BLUE)
    arrow_right(s, Inches(4.3), Inches(2.85), Inches(0.5), Inches(0.4))
    card(s, Inches(4.9), Inches(2.0), Inches(3.5), Inches(2.0),
         "LLM Output", '"metric OR evaluation\nOR BLEU OR ROUGE"', ORANGE)
    arrow_right(s, Inches(8.5), Inches(2.85), Inches(0.5), Inches(0.4))
    card(s, Inches(9.1), Inches(2.0), Inches(3.5), Inches(2.0),
         "Microsoft Graph", "POST /search/query\n-> list of files + snippets", GREEN)

    textbox(s, Inches(0.5), Inches(4.7), Inches(12.5), Inches(0.5),
            "Microsoft has already indexed full text - no need to build it ourselves",
            size=16, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)
    textbox(s, Inches(0.5), Inches(5.5), Inches(12.5), Inches(0.5),
            "Each file comes with a snippet of ~240 characters around the matching keyword",
            size=14, color=GRAY, align=PP_ALIGN.CENTER)

    # ===== 8. Fetch & Extract =====
    s = blank_slide(prs)
    title_bar(s, "Fetch step - download files and extract text", SW)

    textbox(s, Inches(0.5), Inches(1.2), Inches(12.5), Inches(0.5),
            "Only when the LLM judges that snippets are not enough to answer",
            size=16, color=GRAY, align=PP_ALIGN.CENTER)

    # Input formats column
    card(s, Inches(0.7), Inches(2.2), Inches(2.6), Inches(0.9), "PDF", "pdfplumber", BLUE)
    card(s, Inches(0.7), Inches(3.3), Inches(2.6), Inches(0.9), "DOCX", "python-docx", PURPLE)
    card(s, Inches(0.7), Inches(4.4), Inches(2.6), Inches(0.9), "XLSX", "openpyxl", GREEN)
    card(s, Inches(0.7), Inches(5.5), Inches(2.6), Inches(0.9), "PPTX", "python-pptx", ORANGE)

    arrow_right(s, Inches(3.5), Inches(3.7), Inches(0.6), Inches(0.5))

    # Process
    proc = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(4.3), Inches(2.5), Inches(4.5), Inches(3.0))
    fill(proc, LIGHT_BLUE)
    no_line(proc)
    tf = proc.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.25)
    set_text(tf, "Extract Text", size=20, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)
    add_para(tf, "", size=10)
    add_para(tf, "-> Bytes into RAM via io.BytesIO", size=13, color=DARK, align=PP_ALIGN.CENTER)
    add_para(tf, "-> Parse according to each format", size=13, color=DARK, align=PP_ALIGN.CENTER)
    add_para(tf, "-> Markers [PAGE N] / [SHEET] / [SLIDE]", size=13, color=DARK, align=PP_ALIGN.CENTER)
    add_para(tf, "-> Plain text, line by line", size=13, color=DARK, align=PP_ALIGN.CENTER)

    arrow_right(s, Inches(9.0), Inches(3.7), Inches(0.6), Inches(0.5))

    # Cache
    card(s, Inches(9.8), Inches(2.5), Inches(3.0), Inches(3.0),
         "RAM Cache",
         "item_id -> text\n\nSubsequent grep calls\nskip the download",
         TEAL)

    # ===== 9. Grep =====
    s = blank_slide(prs)
    title_bar(s, "Grep step - extract context", SW)

    textbox(s, Inches(0.5), Inches(1.1), Inches(12.5), Inches(0.5),
            "Multi-pattern regex over the cached text",
            size=16, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)

    # Input pattern
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Inches(0.7), Inches(1.9), Inches(5.5), Inches(1.2))
    fill(box, ORANGE)
    no_line(box)
    set_text(box.text_frame, "LLM produces a pattern", size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_para(box.text_frame, "(metric|BLEU|ROUGE|recall|F1)", size=12, color=WHITE, align=PP_ALIGN.CENTER)

    arrow_down(s, Inches(3.0), Inches(3.2), Inches(0.5), Inches(0.4))

    # Match + window
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Inches(0.7), Inches(3.7), Inches(5.5), Inches(3.3))
    fill(box, LIGHT_BG)
    no_line(box)
    tf = box.text_frame
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.15)
    tf.word_wrap = True
    set_text(tf, "Output - chunk +/-5 lines", size=14, bold=True, color=DARK_BLUE)
    add_para(tf, "   L11: Section 4. Evaluation", size=11, color=GRAY)
    add_para(tf, "   L12: ", size=11, color=GRAY)
    add_para(tf, ">> L13: We use BLEU as primary metric", size=11, bold=True, color=BLUE)
    add_para(tf, "   L14: and ROUGE-L for generation.", size=11, color=GRAY)
    add_para(tf, "   L15: ...", size=11, color=GRAY)

    # Right side: features
    card(s, Inches(6.8), Inches(1.9), Inches(6.0), Inches(1.1),
         "Merge overlapping windows",
         "Two nearby hits -> a single chunk, avoiding duplication",
         GREEN)
    card(s, Inches(6.8), Inches(3.1), Inches(6.0), Inches(1.1),
         "Cap at 10 chunks max",
         "Avoid consuming the LLM's entire context window",
         PURPLE)
    card(s, Inches(6.8), Inches(4.3), Inches(6.0), Inches(1.1),
         "Log to debug/grep.log",
         "All context is saved for inspection",
         TEAL)
    card(s, Inches(6.8), Inches(5.5), Inches(6.0), Inches(1.1),
         "Cite line number",
         "The LLM can reference the exact line in its answer",
         BLUE)

    # ===== 10. Final Answer =====
    s = blank_slide(prs)
    title_bar(s, "Final step - synthesis and citations", SW)

    # Input observations
    card(s, Inches(0.7), Inches(1.6), Inches(3.8), Inches(4.5),
         "Inputs to the LLM",
         "- User query\n\n- Search snippets\n\n- Grep chunks\n\n- File metadata\n\n(No full text)",
         PURPLE)

    arrow_right(s, Inches(4.7), Inches(3.6), Inches(0.5), Inches(0.5))

    # LLM brain
    box = s.shapes.add_shape(MSO_SHAPE.OVAL,
                              Inches(5.3), Inches(2.6), Inches(2.8), Inches(2.8))
    fill(box, ORANGE)
    no_line(box)
    tf = box.text_frame
    tf.word_wrap = True
    set_text(tf, "LLM", size=48, color=WHITE, align=PP_ALIGN.CENTER)
    add_para(tf, "LLM synthesizes", size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    arrow_right(s, Inches(8.3), Inches(3.6), Inches(0.5), Inches(0.5))

    # Answer
    card(s, Inches(8.9), Inches(1.6), Inches(3.9), Inches(4.5),
         "Final Answer",
         "Concise answer\n\n+ Citation:\n[file.pdf, L42]\n[doc.docx, PAGE 3]\n\nUser clicks\nto open the source file on OneDrive",
         GREEN)

    textbox(s, Inches(0.5), Inches(6.4), Inches(12.5), Inches(0.6),
            "The LLM must cite - if nothing is found it must say so, never fabricate",
            size=14, color=DARK_BLUE, align=PP_ALIGN.CENTER)

    # ===== 11. Comparison =====
    s = blank_slide(prs)
    title_bar(s, "Comparison - traditional RAG vs this Agent", SW)

    headers = ["Criterion", "Traditional RAG", "Agentic Keyword"]
    rows = [
        ["Infrastructure", "Separate Vector DB", "Not required"],
        ["Chunking", "Mandatory", "Not required"],
        ["Exact match", "May miss", "Perfect"],
        ["Updates", "Re-index everything", "Instant (Microsoft index)"],
        ["Snippet", "Fixed chunks", "Dynamic context +/-N lines"],
        ["Cost", "Embedding + storage", "LLM calls only"],
    ]
    table = s.shapes.add_table(len(rows) + 1, 3,
                                Inches(0.7), Inches(1.3),
                                Inches(11.9), Inches(5.4)).table
    table.columns[0].width = Inches(3.0)
    table.columns[1].width = Inches(4.45)
    table.columns[2].width = Inches(4.45)

    # Header row
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        set_text(cell.text_frame, h, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Data rows
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            cell = table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT_BG if i % 2 == 0 else WHITE
            color = DARK if j == 0 else (ORANGE if j == 1 else GREEN)
            set_text(cell.text_frame, v, size=14,
                     bold=(j == 0), color=color, align=PP_ALIGN.LEFT)

    # ===== 12. Conclusion =====
    s = blank_slide(prs)
    title_bar(s, "Conclusion", SW)

    points = [
        ("*", "Simpler", "No vector DB, no re-indexing"),
        (">", "Instant", "Microsoft already indexes - queries are ready immediately"),
        ("o", "Exact match", "Never miss product codes / technical terms"),
        ("A", "LLM-driven", "Generates keywords, retries strategies, cites sources"),
        ("=", "Transparent", "Every step streams in realtime with full logging"),
    ]
    for i, (ico, h, t) in enumerate(points):
        y = Inches(1.4 + i * 1.05)
        ic = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), y, Inches(0.8), Inches(0.8))
        fill(ic, BLUE)
        no_line(ic)
        set_text(ic.text_frame, ico, size=22, color=WHITE, align=PP_ALIGN.CENTER)

        textbox(s, Inches(2.0), y + Inches(0.05), Inches(3.0), Inches(0.5),
                h, size=18, bold=True, color=DARK_BLUE)
        textbox(s, Inches(5.2), y + Inches(0.1), Inches(7.5), Inches(0.5),
                t, size=14, color=GRAY)

    textbox(s, Inches(0.5), Inches(6.9), Inches(12.5), Inches(0.5),
            "Microsoft 365 Keyword-Search Agent",
            size=12, bold=True, color=BLUE, align=PP_ALIGN.CENTER)

    prs.save(path)
    print(f"[OK] Saved {path} ({len(prs.slides)} slides)")


# ---------- docx helpers ----------
def add_heading(doc, text, level=1, color=None):
    h = doc.add_heading(text, level=level)
    if color is not None:
        for run in h.runs:
            run.font.color.rgb = color
    return h


def add_body(doc, text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = DocxPt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = "Calibri"
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    if level:
        p.paragraph_format.left_indent = DocxPt(18 * (level + 1))
    for run in p.runs:
        run.font.size = DocxPt(11)
        run.font.name = "Calibri"
    return p


# ---------- docx generation ----------
def make_docx(path):
    doc = Document()

    # ====== Title ======
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Microsoft 365 Keyword-Search Agent")
    r.font.size = DocxPt(24)
    r.font.bold = True
    r.font.color.rgb = DocxRGB(0x00, 0x78, 0xD4)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Workflow report")
    r.font.size = DocxPt(14)
    r.font.italic = True
    r.font.color.rgb = DocxRGB(0x60, 0x5E, 0x5C)

    doc.add_paragraph()

    # ====== 1. Introduction ======
    add_heading(doc, "1. Introduction", level=1, color=DocxRGB(0x00, 0x78, 0xD4))
    add_body(doc,
        "A web application that lets users ask questions about OneDrive "
        "documents in natural language. Inspired by the paper \"Keyword "
        "search is all you need: Achieving RAG-Level Performance without "
        "vector databases using agentic tool use\" (arxiv 2602.23368), the "
        "system replaces the traditional Retrieval-Augmented Generation "
        "(RAG) architecture with an intelligent LLM agent that uses the "
        "keyword search already provided by the Microsoft Graph API.")

    # ====== 2. Problem ======
    add_heading(doc, "2. The problem with traditional RAG", level=1, color=DocxRGB(0x00, 0x78, 0xD4))
    add_body(doc, "Traditional RAG requires a complex processing pipeline:")
    for b in [
        "Chunking - splitting documents into smaller pieces. Easy to lose context when cutting mid-sentence or through a table.",
        "Embedding - creating a vector for each chunk. Costly in compute and storage.",
        "Vector database - operating an additional service (FAISS, Pinecone, Weaviate, etc.).",
        "Re-indexing - re-creating embeddings is expensive whenever documents change.",
        "Weak exact match - semantic similarity can miss product codes and rare technical terms.",
    ]:
        add_bullet(doc, b)

    # ====== 3. Solution ======
    add_heading(doc, "3. Proposed solution", level=1, color=DocxRGB(0x00, 0x78, 0xD4))
    add_body(doc,
        "Instead of building a RAG pipeline, the system uses an LLM agent "
        "that drives the search by itself. The core ideas are:")
    for b in [
        "Microsoft already provides a full-text index for OneDrive/SharePoint - leverage it via the Graph Search API.",
        "The LLM generates many keyword variants (synonyms, abbreviations, regex) - compensating for the index being literal-only.",
        "Download files and extract text only when snippets are not enough - keep the original context, no pre-chunking.",
        "Regex grep with a +/-N line context window - simulating a \"dynamic chunk\" placed on actual hits.",
        "The LLM evaluates results itself and retries with a different strategy if nothing is found.",
    ]:
        add_bullet(doc, b)

    # ====== 4. Actors ======
    add_heading(doc, "4. Participating actors", level=1, color=DocxRGB(0x00, 0x78, 0xD4))

    add_heading(doc, "4.1. User", level=2)
    add_body(doc, "Signs in with a Microsoft 365 account and types natural-language questions into the chat panel on the web dashboard.")

    add_heading(doc, "4.2. Web UI (Browser)", level=2)
    add_body(doc,
        "The dashboard page displays OneDrive/Teams/SharePoint information for the user. "
        "A JavaScript chat panel sends requests to the server and displays each agent step "
        "in real time via Server-Sent Events (SSE).")

    add_heading(doc, "4.3. Flask Backend", level=2)
    add_body(doc,
        "A Python server that handles OAuth2 sign-in (Microsoft Entra ID), manages "
        "sessions, and exposes the /ask route to start the agent. It passes the "
        "access_token to the agent so it can call Microsoft Graph on behalf of the user.")

    add_heading(doc, "4.4. LangChain Agent (langgraph)", level=2)
    add_body(doc,
        "The tool-use loop coordinator. It gives the question and system prompt to the LLM, "
        "receives a decision about which tool to call, runs that tool, and feeds the result "
        "back to the LLM. It loops up to 30 times until the LLM produces a final answer.")

    add_heading(doc, "4.5. LLM (Azure OpenAI)", level=2)
    add_body(doc,
        "The decision-making brain. It reads the question and interaction history to decide: "
        "which tool to call and with which arguments, or whether it already has enough to answer. "
        "It generates OR-joined keyword variants for search and multi-pattern regex for grep. "
        "Default model: gpt-4.1-mini.")

    add_heading(doc, "4.6. Microsoft Graph API", level=2)
    add_body(doc,
        "Microsoft's official endpoint for 365 data. Two endpoints matter:")
    for b in [
        "POST /search/query - full-text search over OneDrive, SharePoint, and Teams files. Returns metadata plus a snippet around the keyword.",
        "GET /drives/{driveId}/items/{id}/content - downloads the binary content of a file.",
    ]:
        add_bullet(doc, b)

    add_heading(doc, "4.7. The Agent's three tools", level=2)
    add_body(doc, "Each tool is a Python function registered with LangChain:")
    for b in [
        "graph_search(query) - calls Microsoft Graph Search and returns a list of files matching the keyword.",
        "fetch_file_text(item_id) - downloads the binary file, extracts text, and caches it in RAM.",
        "grep_context(item_id, patterns, context_lines) - regex-searches the cached text and returns chunks of +/-N lines around each hit.",
    ]:
        add_bullet(doc, b)

    add_heading(doc, "4.8. File Extractors", level=2)
    add_body(doc, "Four modules that turn binary content into plain text, line by line:")
    for b in [
        "PDF - uses pdfplumber, marks [PAGE N] for each page.",
        "DOCX - uses python-docx, captures paragraphs and tables.",
        "XLSX - uses openpyxl, marks [SHEET name] per sheet, joins each row with a pipe.",
        "PPTX - uses python-pptx, marks [SLIDE N], including speaker notes.",
    ]:
        add_bullet(doc, b)

    add_heading(doc, "4.9. Cache and Drive Index", level=2)
    add_body(doc,
        "Two in-memory dicts that optimize performance:")
    for b in [
        "Text cache (cache.py) - once a file is extracted, the text is stored in RAM. Later grep calls do not need to re-download.",
        "Drive index (drive_index.py) - stores the mapping item_id -> driveId, so fetch uses the right endpoint for shared / SharePoint / Teams files.",
    ]:
        add_bullet(doc, b)

    # ====== 5. End-to-end workflow ======
    add_heading(doc, "5. End-to-end workflow", level=1, color=DocxRGB(0x00, 0x78, 0xD4))

    steps = [
        ("Step 1 - User enters a question",
         "The user types a natural-language question into the dashboard chat panel, for example "
         "\"What is the metrics for RAG?\" and hits Send."),
        ("Step 2 - Browser sends the request",
         "JavaScript intercepts the form submit and calls fetch POST /ask with JSON "
         "containing the question. It accepts a text/event-stream response."),
        ("Step 3 - Flask initializes the agent",
         "The server takes the access_token from the existing session (from the OAuth2 "
         "flow with MSAL). It initializes the LangChain agent with the three tools, "
         "each bound to the token via a closure. It opens the SSE stream."),
        ("Step 4 - LLM generates search keywords",
         "The LLM reads the question and the system prompt. The system prompt requires the LLM "
         "to use multi-pattern OR queries (never search for a single word). "
         "The LLM emits a graph_search tool call with a query like \"metric OR "
         "evaluation OR BLEU OR ROUGE OR recall\"."),
        ("Step 5 - Microsoft Graph finds files",
         "The graph_search tool calls POST /search/query with entityTypes "
         "[driveItem]. The Microsoft index returns up to 15 matching files, "
         "each with a snippet of about 240 characters around the matching keyword."),
        ("Step 6 - LLM evaluates the snippets",
         "The LLM receives the list of files plus snippets. Two possibilities: "
         "(a) the snippets are enough to answer - jump to Step 9; "
         "(b) deeper reading is required - move to Step 7."),
        ("Step 7 - Download and extract files",
         "The LLM calls fetch_file_text for the 1-3 most relevant files. The tool looks up "
         "drive_index to find the driveId, downloads the file via the Graph endpoint "
         "(which auto-redirects to Microsoft's CDN), extracts the text with the matching "
         "extractor, and stores it in the cache. The LLM receives a short preview so the "
         "context window is not exhausted."),
        ("Step 8 - Grep context",
         "The LLM calls grep_context with multi-pattern regex and context_lines=5. "
         "The tool finds matches over the cached text, merges overlapping windows, "
         "and returns up to 10 chunks. Each chunk uses \">>\" for hit lines and "
         "\"  \" for context lines. It also logs everything to debug/grep.log."),
        ("Step 9 - LLM synthesizes the answer",
         "Once it has enough context, the LLM produces the final answer with "
         "citations in the form [filename, L42] or [filename, PAGE 3]. "
         "If nothing is found, the LLM must say so explicitly - "
         "fabrication is not allowed."),
        ("Step 10 - Stream back to the browser",
         "Every step (tool call, tool result, final answer, summary) is "
         "yielded as an SSE event. The browser receives and renders each event "
         "in real time on the chat panel."),
    ]
    for h, t in steps:
        add_heading(doc, h, level=3)
        add_body(doc, t)

    # ====== 6. SSE ======
    add_heading(doc, "6. Server-Sent Events (SSE)", level=1, color=DocxRGB(0x00, 0x78, 0xD4))
    add_body(doc,
        "SSE lets the server continuously push events to the browser over a "
        "single HTTP connection. In this application, every agent step "
        "(LLM decides to call a tool, tool returns a result, LLM generates the "
        "final answer) is sent as soon as it happens. The browser displays it "
        "live so the user can see \"what the agent is thinking\" without having "
        "to wait for the entire process to finish.")

    # ====== 7. Cache ======
    add_heading(doc, "7. The role of cache and drive index", level=1, color=DocxRGB(0x00, 0x78, 0xD4))
    add_body(doc,
        "Each file is downloaded only once. After extraction, the text is "
        "stored in an in-memory dict keyed by the Graph item_id. Subsequent "
        "grep_context calls read directly from RAM without any network call. "
        "The drive index stores the mapping from item_id to driveId because "
        "files in OneDrive are not always under /me/drive - they may be "
        "shared files, from a SharePoint site, or from a Teams group drive.")

    # ====== 8. Comparison ======
    add_heading(doc, "8. Comparison with traditional RAG", level=1, color=DocxRGB(0x00, 0x78, 0xD4))
    table = doc.add_table(rows=7, cols=3)
    table.style = "Light Grid Accent 1"
    headers = ["Criterion", "Traditional RAG", "Agentic Keyword (this system)"]
    rows = [
        ["Infrastructure", "Separate Vector DB (FAISS, Pinecone, etc.)", "Not required - uses Microsoft's existing index"],
        ["Chunking", "Mandatory, fixed in advance", "Not required - context window is dynamic per hit"],
        ["Exact match", "May miss (semantic similarity)", "Perfect (literal regex)"],
        ["Document updates", "Re-index everything", "Instant (Microsoft indexes automatically)"],
        ["Cost", "Embedding + storage + LLM", "LLM calls only"],
        ["Transparency", "Hard to debug similarity scores", "Every step is fully logged and traceable"],
    ]
    cells = table.rows[0].cells
    for i, h in enumerate(headers):
        cells[i].text = ""
        run = cells[i].paragraphs[0].add_run(h)
        run.font.bold = True
        run.font.size = DocxPt(11)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            table.rows[i].cells[j].text = ""
            run = table.rows[i].cells[j].paragraphs[0].add_run(v)
            run.font.size = DocxPt(10)

    # ====== 9. Limitations ======
    add_heading(doc, "9. Current limitations", level=1, color=DocxRGB(0x00, 0x78, 0xD4))
    for b in [
        "Scanned PDFs (image-only) - pdfplumber cannot OCR, so it returns empty text. Azure Document Intelligence or Docling would be needed for OCR.",
        "25 MB per-file limit - larger files are rejected to avoid RAM pressure and timeouts.",
        "Speed depends on the number of LLM calls - each question usually costs 5-15 calls through the ReAct loop.",
        "Accuracy depends on model quality - gpt-4.1-mini works well; weaker models may not generate enough keyword variants.",
        "Microsoft's search index lag - newly uploaded files may take several minutes to be indexed.",
    ]:
        add_bullet(doc, b)

    # ====== 10. Conclusion ======
    add_heading(doc, "10. Conclusion", level=1, color=DocxRGB(0x00, 0x78, 0xD4))
    add_body(doc,
        "The system demonstrates that with a sufficiently capable LLM and the "
        "right keyword-search toolkit, we can achieve quality comparable to RAG "
        "without building and operating a vector database. By leveraging the "
        "Microsoft Graph Search that already exists and letting the LLM drive "
        "the search process, the architecture becomes simpler, cheaper, easier "
        "to debug, and supports perfect exact match for technical documents.")

    # Footer
    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = foot.add_run("- Microsoft 365 Keyword-Search Agent -")
    r.font.size = DocxPt(10)
    r.font.italic = True
    r.font.color.rgb = DocxRGB(0x60, 0x5E, 0x5C)

    doc.save(path)
    print(f"[OK] Saved {path}")


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    make_pptx(os.path.join(here, "graph_grep.pptx"))
    make_docx(os.path.join(here, "docs_report.docx"))
