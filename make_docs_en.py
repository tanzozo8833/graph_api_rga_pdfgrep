"""Generate docs_report_en.docx -- product report (English).

Run: venv\\Scripts\\python.exe make_docs_en.py
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

DARK = RGBColor(0x20, 0x20, 0x20)
ACCENT = RGBColor(0x00, 0x47, 0x82)
GRAY = RGBColor(0x55, 0x55, 0x55)

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"


def set_run(run, *, size=11, bold=False, italic=False, color=DARK, font=BODY_FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_paragraph(doc, text="", *, size=11, bold=False, italic=False,
                  color=DARK, font=BODY_FONT, align=WD_ALIGN_PARAGRAPH.LEFT,
                  space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, italic=italic, color=color, font=font)
    return p


def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run(r, size=18, bold=True, color=ACCENT)
    return p


def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run(r, size=13, bold=True, color=DARK)
    return p


def add_bullet(doc, text, *, indent=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.6 + 0.5 * indent)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_run(r, size=11)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_after = Pt(6)
    for i, line in enumerate(text.splitlines()):
        if i > 0:
            p.add_run().add_break()
        r = p.add_run(line if line else " ")
        set_run(r, size=10, color=DARK, font=MONO_FONT)
    return p


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{label}: ")
    set_run(r1, size=11, bold=True)
    r2 = p.add_run(value)
    set_run(r2, size=11)
    return p


def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("_" * 80)
    set_run(r, size=8, color=GRAY)


# ---------- Cover ----------

def cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Microsoft 365 Keyword-Search Agent")
    set_run(r, size=26, bold=True, color=ACCENT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(40)
    r = p.add_run("Product Report")
    set_run(r, size=14, italic=True, color=GRAY)

    add_label_value(doc, "Project", "Question answering over OneDrive without a vector database")
    add_label_value(doc, "Stack", "Flask + MSAL + LangChain (langgraph) + Azure OpenAI + Microsoft Graph")
    add_label_value(doc, "Status", "Working prototype")

    doc.add_page_break()


# ---------- 1. Problems define ----------

def section_problems(doc):
    add_heading_1(doc, "1. Problems Define")

    add_paragraph(
        doc,
        "Two real problems motivate this project. They are independent but they show up "
        "together inside almost every enterprise that wants to do question answering over "
        "its own documents."
    )

    add_heading_2(doc, "1.1 Problem 1 - Documents are scattered across the enterprise")
    add_paragraph(
        doc,
        "In a modern Microsoft 365 organisation, an employee's working documents do not "
        "live in a single place. The same project usually has artefacts spread across "
        "several storage surfaces:"
    )
    for line in [
        "OneDrive: personal drafts, individual reports, ad-hoc spreadsheets.",
        "SharePoint: team libraries, official policies, shared product documents.",
        "Microsoft Teams: files attached to channels, meeting notes, recordings.",
        "Outlook attachments: contracts and PDFs that never make it to a shared drive.",
    ]:
        add_bullet(doc, line)
    add_paragraph(
        doc,
        "Building a question-answering tool that crawls each of these surfaces separately "
        "is painful: each one has its own API, its own auth flow, its own permission model "
        "and its own rate limits. The result is brittle and expensive to maintain.",
        space_after=4,
    )
    add_paragraph(doc, "Our approach to Problem 1:", bold=True, space_after=2)
    add_paragraph(
        doc,
        "Use Microsoft Graph as the single entry point. Graph already federates OneDrive, "
        "SharePoint, Teams and Outlook behind one HTTPS endpoint, one OAuth token and one "
        "permission model. POST /search/query returns ranked, snippet-rich results from "
        "all of those surfaces at once - the agent never needs to know which surface a "
        "document came from."
    )

    add_heading_2(doc, "1.2 Problem 2 - Classical RAG is heavy and lossy")
    add_paragraph(
        doc,
        "The default solution for document question answering is Retrieval-Augmented "
        "Generation (RAG): chunk every document, embed every chunk, store the vectors in "
        "a vector database, and query by semantic similarity. This works, but it is "
        "expensive and it loses information:"
    )
    for line in [
        "Chunking breaks structure - tables, code blocks and multi-paragraph arguments "
        "get split across chunk boundaries.",
        "Embedding is a recurring cost - every new or edited document must be re-embedded; "
        "the vector store must be re-indexed.",
        "A vector database is an extra service to deploy, secure and back up.",
        "Embeddings miss exact matches - rare technical terms, product codes, version "
        "numbers and acronyms often score poorly against a paraphrased query.",
        "Single-shot retrieval - classical RAG retrieves once and then answers; if the "
        "first retrieval is weak, the model has no second chance.",
        "Index drift - when a user uploads or deletes a file, the vector store is stale "
        "until the next re-indexing job runs.",
    ]:
        add_bullet(doc, line)
    add_paragraph(doc, "Our approach to Problem 2:", bold=True, space_after=2)
    add_paragraph(
        doc,
        "Replace the RAG pipeline with an agentic loop driven by the LLM. The LLM "
        "generates keywords, calls Microsoft Graph search, reads snippets, and only "
        "downloads a file when the snippet is not enough. Inside a downloaded file, a "
        "simple regex grep with +/- N lines of context recovers the exact passage. No "
        "chunking, no embedding, no vector database."
    )

    add_hr(doc)


# ---------- 2. Solution ----------

def section_solution(doc):
    add_heading_1(doc, "2. Solution")

    add_paragraph(
        doc,
        "The solution is one Flask web app that combines two ideas:"
    )
    for line in [
        "Microsoft Graph as the universal document API across OneDrive, SharePoint, Teams "
        "and Outlook attachments.",
        "An LLM agent that searches by keyword and refines its answer with grep, instead "
        "of pre-indexing everything into a vector database.",
    ]:
        add_bullet(doc, line)

    add_heading_2(doc, "2.1 Unified access via Microsoft Graph")
    add_paragraph(
        doc,
        "The user signs in once with their Microsoft 365 account using the MSAL "
        "Authorization Code flow. The resulting access token grants the agent read access "
        "to every document the user can already see - across OneDrive, SharePoint sites, "
        "Teams channel files and supported Outlook attachments - through one endpoint: "
        "POST /search/query."
    )
    for line in [
        "One auth flow instead of four.",
        "One ranked result set instead of four parallel searches to merge.",
        "Permissions inherited from the signed-in user - the agent cannot see anything the "
        "user could not see directly.",
    ]:
        add_bullet(doc, line)

    add_heading_2(doc, "2.2 Agentic keyword search instead of RAG")
    add_paragraph(
        doc,
        "Inspired by the paper \"Keyword search is all you need\" (arXiv 2602.23368), the "
        "LLM drives the search itself. It generates OR-joined keyword sets (synonyms, "
        "abbreviations, numeric variants), reads the snippets returned by Graph, and "
        "decides whether to fetch a file and grep inside it."
    )

    add_heading_2(doc, "2.3 The three tools given to the LLM")
    for label, body in [
        ("graph_search(query)",
         "Calls POST /search/query on Microsoft Graph. Returns matching files (from "
         "OneDrive, SharePoint, Teams, ...) with a short snippet around the keyword. "
         "Uses OR-joined patterns for recall."),
        ("fetch_file_text(item_id)",
         "Downloads a file via Graph and extracts plain text using pdfplumber / "
         "python-docx / openpyxl / python-pptx. The extracted text is cached in memory "
         "for the rest of the conversation."),
        ("grep_context(item_id, patterns, context_lines)",
         "Regex search over the cached text with N lines of surrounding context. Supports "
         "multi-pattern OR, merges overlapping windows and reports line numbers used in "
         "the final citation."),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(label + " - ")
        set_run(r1, size=11, bold=True, font=MONO_FONT)
        r2 = p.add_run(body)
        set_run(r2, size=11)

    add_heading_2(doc, "2.4 Why this is different from classical RAG")
    for line in [
        "No pre-processing pipeline: nothing is chunked, embedded or stored ahead of time.",
        "Live data: every query hits Microsoft 365 directly, so newly uploaded or edited "
        "files are immediately searchable.",
        "Transparent: each step (keyword used, file fetched, regex matched) is streamed to "
        "the UI through Server-Sent Events, so the agent's reasoning is fully auditable.",
        "Cheaper to operate: no vector database to host, no re-embedding job to schedule, "
        "no embedding API bill.",
    ]:
        add_bullet(doc, line)

    add_hr(doc)


# ---------- 3. Goal of this research ----------

def section_goal(doc):
    add_heading_1(doc, "3. Goal of This Research")

    add_paragraph(
        doc,
        "This project sets out to validate, with a working prototype, that a single LLM "
        "agent on top of Microsoft Graph can replace a traditional multi-connector RAG "
        "stack for enterprise document question answering."
    )

    add_heading_2(doc, "3.1 Primary goals")
    for line in [
        "Aggregate documents from multiple sources through one API: OneDrive, SharePoint, "
        "Microsoft Teams files and Outlook attachments are all reached via Microsoft Graph "
        "with a single token.",
        "Replace vector-based RAG with fast keyword search and grep: rely on Microsoft's "
        "existing full-text index for the first pass, and use regex with surrounding "
        "context to extract exact passages.",
        "Provide cited answers: every final answer must include a [filename] or "
        "[filename, line N] citation pointing back to the source document.",
        "Keep the operational footprint minimal: no vector store, no background indexer, "
        "no embedding API cost - just a Flask process and two SaaS calls (Graph, Azure "
        "OpenAI).",
    ]:
        add_bullet(doc, line)

    add_heading_2(doc, "3.2 Secondary goals")
    for line in [
        "Make the reasoning visible: each search keyword, file fetch and regex hit is "
        "streamed to the browser in real time over Server-Sent Events.",
        "Respect Microsoft 365 permissions: the agent inherits the signed-in user's "
        "access rights and cannot see anything the user could not see directly.",
        "Keep the codebase small and readable: each tool is a single file, the agent loop "
        "is a few hundred lines, no hidden orchestration framework.",
        "Support common file formats out of the box: PDF, DOCX, XLSX, PPTX and plain text.",
    ]:
        add_bullet(doc, line)

    add_heading_2(doc, "3.3 Out of scope for this prototype")
    for line in [
        "OCR of scanned (image-only) PDFs.",
        "Files larger than 25 MB.",
        "Cross-tenant search or anonymous access.",
        "A persistent storage layer (the file-text cache lives only in process memory).",
    ]:
        add_bullet(doc, line)

    add_hr(doc)


# ---------- 4. Summarize conclusion ----------

def section_conclusion(doc):
    add_heading_1(doc, "4. Summarize Conclusion")

    add_paragraph(
        doc,
        "The prototype answers the two motivating problems with one design. Microsoft "
        "Graph solves the document-scatter problem, and an LLM agent with keyword search "
        "plus grep solves the RAG-overhead problem. The combined system is simpler, "
        "cheaper to run and easier to audit than the classical multi-connector RAG stack "
        "it replaces."
    )

    add_heading_2(doc, "4.1 Conclusions on Problem 1 (scattered documents)")
    for line in [
        "Microsoft Graph is sufficient as a single front door to OneDrive, SharePoint, "
        "Teams files and Outlook attachments - one token, one search endpoint, one ranked "
        "result set.",
        "Permission handling is solved for free: every call inherits the signed-in user's "
        "rights, so the agent cannot expose data the user could not already access.",
        "New or edited documents are visible to the agent immediately - no re-indexing "
        "job, no drift between source and search.",
    ]:
        add_bullet(doc, line)

    add_heading_2(doc, "4.2 Conclusions on Problem 2 (RAG overhead)")
    for line in [
        "Keyword search plus grep is enough for typical enterprise questions: snippets "
        "from Microsoft Graph answer many queries directly; for the rest, a regex pass "
        "with +/- 5 lines of context finds the exact passage.",
        "The LLM does the heavy lifting that embeddings used to do: it generates OR-joined "
        "keyword sets, judges when a snippet is sufficient, and retries with new keyword "
        "strategies when the first attempt is weak.",
        "Exact matches (product codes, version numbers, acronyms) are recovered reliably, "
        "because the system searches the literal text instead of an approximate vector.",
        "Operating cost drops sharply: no vector database, no embedding API bill, no "
        "scheduled re-indexing - just one Flask process, an in-memory cache and two SaaS "
        "endpoints.",
    ]:
        add_bullet(doc, line)

    add_heading_2(doc, "4.3 Trade-offs to be aware of")
    for line in [
        "Latency depends on Microsoft Graph: a search is one network round-trip; fetching "
        "a large PDF can take a few seconds.",
        "Quality depends on keyword coverage: when the user's wording is far from the "
        "wording in the document, the LLM has to try several keyword strategies.",
        "Cache lifetime: the in-memory text cache is lost on restart. Acceptable for a "
        "single-user prototype but would need a real store for multi-user production.",
        "Pure semantic queries (paraphrase, cross-language matching, summarisation across "
        "many documents) remain a better fit for vector-based RAG.",
    ]:
        add_bullet(doc, line)

    add_heading_2(doc, "4.4 Recommendation")
    add_paragraph(
        doc,
        "For enterprise document question answering on Microsoft 365 content, ship this "
        "agent first. It removes an entire vector-database tier, unifies four storage "
        "surfaces behind one API and produces cited answers out of the box. Add a "
        "vector-based layer later, only for the specific queries that keyword search "
        "demonstrably cannot serve."
    )

    add_hr(doc)


# ---------- 5. Detail about technical ----------

def section_technical(doc):
    add_heading_1(doc, "5. Detail About Technical")

    # 5.1 Setup environment / prerequisites
    add_heading_2(doc, "5.1 Setup Environment and Prerequisites")

    add_paragraph(doc, "Software prerequisites:", bold=True, space_after=2)
    for line in [
        "Python 3.10 or newer (tested on 3.11 / 3.12).",
        "A Windows, macOS or Linux developer machine.",
        "A Microsoft 365 account with access to OneDrive.",
        "An Azure subscription with access to Azure OpenAI.",
    ]:
        add_bullet(doc, line)

    add_paragraph(doc, "Azure App Registration:", bold=True, space_after=2)
    for line in [
        "Register an application in Microsoft Entra ID (Azure Active Directory).",
        "Redirect URI: http://localhost:3000/auth/callback (type: Web).",
        "API permissions (delegated): User.Read, Mail.Read, Files.Read, Files.Read.All, "
        "Sites.Read.All, Team.ReadBasic.All.",
        "Create a client secret under Certificates & secrets and copy the value to .env.",
    ]:
        add_bullet(doc, line)

    add_paragraph(doc, "Azure OpenAI:", bold=True, space_after=2)
    for line in [
        "Provision an Azure OpenAI resource and deploy a tool-calling model (e.g. gpt-4o "
        "or gpt-4.1-mini).",
        "Note the endpoint, API key, deployment name and API version.",
    ]:
        add_bullet(doc, line)

    add_paragraph(doc, "Install dependencies:", bold=True, space_after=2)
    add_code(doc,
             "python -m venv venv\n"
             "venv\\Scripts\\activate          # Windows\n"
             "source venv/bin/activate         # macOS / Linux\n"
             "pip install -r requirements.txt")

    add_paragraph(doc, "Environment variables (.env file):", bold=True, space_after=2)
    add_code(doc,
             "CLIENT_ID=...\n"
             "TENANT_ID=...\n"
             "CLIENT_SECRET=...\n"
             "FLASK_SECRET_KEY=any-random-string\n"
             "\n"
             "AZURE_OPENAI_API_KEY=...\n"
             "AZURE_OPENAI_ENDPOINT=https://<resource>.openai.azure.com/\n"
             "AZURE_OPENAI_DEPLOYMENT=gpt-4o\n"
             "AZURE_OPENAI_API_VERSION=2024-10-21")

    add_paragraph(doc, "Verify the setup:", bold=True, space_after=2)
    add_code(doc,
             "python test_azure.py   # smoke test for Azure OpenAI credentials\n"
             "python app.py         # start the Flask app on http://localhost:3000")

    # 5.2 Processing flows
    add_heading_2(doc, "5.2 Processing Flows")

    add_paragraph(doc, "End-to-end request flow:", bold=True, space_after=2)
    add_code(doc,
             "1. User signs in at http://localhost:3000 (MSAL Authorization Code flow).\n"
             "2. User types a question in the chat panel.\n"
             "3. Browser opens an SSE connection to /ask.\n"
             "4. Flask creates a LangChain agent (langgraph) bound to the user's access token.\n"
             "5. Agent calls graph_search with OR-joined keywords.\n"
             "6. Agent reads snippets:\n"
             "     - if enough -> produce Final Answer with citation\n"
             "     - if partial -> fetch_file_text on the most relevant files\n"
             "7. Agent calls grep_context to extract the exact passages.\n"
             "8. Agent either answers, or retries with a new keyword strategy.\n"
             "9. Every step is streamed to the browser as an SSE event.")

    add_paragraph(doc, "Agent decision loop (per turn):", bold=True, space_after=2)
    add_code(doc,
             "      +---------------------------+\n"
             "      |  User question (NL)       |\n"
             "      +-------------+-------------+\n"
             "                    |\n"
             "                    v\n"
             "      +---------------------------+\n"
             "      |  LLM generates keywords   |\n"
             "      |  (synonyms, OR-joined)    |\n"
             "      +-------------+-------------+\n"
             "                    |\n"
             "                    v\n"
             "      +---------------------------+\n"
             "      |  graph_search(query)      |--> snippets enough? --> Final Answer\n"
             "      +-------------+-------------+\n"
             "                    | no\n"
             "                    v\n"
             "      +---------------------------+\n"
             "      |  fetch_file_text(item_id) |\n"
             "      +-------------+-------------+\n"
             "                    |\n"
             "                    v\n"
             "      +---------------------------+\n"
             "      |  grep_context(regex, +/-N)|--> found?  yes --> Final Answer + citation\n"
             "      +-------------+-------------+         no\n"
             "                    |\n"
             "                    v\n"
             "          retry with new keywords")

    add_paragraph(doc, "Data flow and caching:", bold=True, space_after=2)
    for line in [
        "graph_search results: not cached. Each call hits Microsoft Graph live so the "
        "agent always sees fresh search rankings.",
        "fetch_file_text: extracted plain text is cached in-process keyed by item_id. "
        "Calling fetch a second time for the same item is a no-op.",
        "grep_context: stateless. Operates entirely on the cached text.",
        "Cache scope: per Flask process. Cleared on restart. No persistence on disk.",
    ]:
        add_bullet(doc, line)

    add_paragraph(doc, "Streaming protocol:", bold=True, space_after=2)
    for line in [
        "Transport: Server-Sent Events on /ask.",
        "Event types: tool_call, tool_result, final_answer, summary, error.",
        "Each event carries a step number and timestamp so the UI can render the "
        "reasoning trace in order.",
    ]:
        add_bullet(doc, line)

    add_hr(doc)


# ---------- 6. Remaining problems ----------

def section_remaining(doc):
    add_heading_1(doc, "6. Remaining Problems")

    add_paragraph(
        doc,
        "The prototype works on text-based documents, but one important scenario has not "
        "been validated yet: files whose informational content is carried by images rather "
        "than by digital text."
    )

    add_heading_2(doc, "6.1 Image-bearing files and Microsoft's OCR")
    add_paragraph(
        doc,
        "Microsoft 365 runs server-side OCR on supported image and scanned-document files "
        "(image attachments, scanned PDFs, image-only slides) and merges the extracted "
        "text into its full-text search index. In theory, this means that POST "
        "/search/query should return such files when the user searches for words that "
        "only appear inside the image."
    )
    for line in [
        "We have not yet measured how reliable that OCR coverage is in practice across "
        "OneDrive, SharePoint and Teams.",
        "We have not measured ranking quality - whether an OCR-extracted hit competes "
        "fairly with the same word found in native text.",
        "It is unclear how quickly newly uploaded image files become searchable (OCR is "
        "typically asynchronous and may take minutes to hours).",
        "Snippet quality for OCR'd content is unknown - the snippet around the matched "
        "keyword may be truncated or noisy.",
    ]:
        add_bullet(doc, line)

    add_heading_2(doc, "6.2 Downstream impact on the agent")
    for line in [
        "graph_search may return an image-only file with a useful snippet, but "
        "fetch_file_text would still extract nothing locally (pdfplumber / python-docx do "
        "not OCR).",
        "In that case the agent has to either trust the Graph-side snippet, or skip the "
        "file - we need to decide and document the policy.",
        "If Microsoft's OCR text is not retrievable through the file download API, the "
        "agent loses the ability to grep for surrounding context.",
    ]:
        add_bullet(doc, line)

    add_heading_2(doc, "6.3 What needs to be tested")
    for line in [
        "Upload a representative set of image files (PNG / JPG / scanned PDF / image-only "
        "PPTX) to OneDrive, SharePoint and Teams.",
        "Wait long enough for Microsoft 365 to OCR them, then issue queries containing "
        "words that only appear inside the images.",
        "Measure: recall (does the file appear in search results?), latency to indexing, "
        "snippet quality, and whether the OCR text is accessible through the file content "
        "endpoint.",
        "Decide the agent policy for image-only hits: cite directly from the Graph "
        "snippet, fall back to a different tool, or mark them as unsupported.",
    ]:
        add_bullet(doc, line)

    add_paragraph(
        doc,
        "Closing this gap is the main follow-up item before the prototype can be claimed "
        "to cover the full range of enterprise documents."
    )

    add_hr(doc)


# ---------- Build ----------

def build():
    doc = Document()

    # Default style
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(11)

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    cover(doc)
    section_problems(doc)
    section_solution(doc)
    section_goal(doc)
    section_conclusion(doc)
    section_technical(doc)
    section_remaining(doc)

    out = "docs_report_en.docx"
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    build()
