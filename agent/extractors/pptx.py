import io

from pptx import Presentation


def extract_pptx(content: bytes) -> str:
    prs = Presentation(io.BytesIO(content))
    lines = []
    for i, slide in enumerate(prs.slides, start=1):
        lines.append(f"[SLIDE {i}]")
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if text:
                    lines.append(text)
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None:
                notes = notes_frame.text.strip()
                if notes:
                    lines.append(f"[NOTES] {notes}")
    return "\n".join(lines)
