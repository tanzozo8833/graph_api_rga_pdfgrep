import io

from docx import Document


def extract_docx(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    for t_idx, table in enumerate(doc.tables, start=1):
        lines.append(f"[TABLE {t_idx}]")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            lines.append(" | ".join(cells))
    return "\n".join(lines)
