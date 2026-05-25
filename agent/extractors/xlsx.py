import io

from openpyxl import load_workbook


def extract_xlsx(content: bytes) -> str:
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"[SHEET {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            cells = ["" if c is None else str(c).replace("\n", " ") for c in row]
            lines.append(" | ".join(cells))
    return "\n".join(lines)
